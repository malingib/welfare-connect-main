from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


VILLAGES = [
    "Malanga – Kabiranduni", "Malanga – Chembe", "Malanga – Kibaoni",
    "Malanga – Ziani", "Malanga – Soyosoyo", "Malanga – Muthoroni",
    "Malanga – Yembe", "Malanga – Majengo", "Malanga – Ngamani",
    "Malanga – Kadzitosoni", "Malanga – Kisimani", "Malanga – Bahati",
    "Malanga – Muungano", "Malanga – Malanga",
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    p.paragraph_format.space_after = Pt(0)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def setup(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for name, size, color in [("Title", 20, "1F4E79"), ("Heading 1", 15, "1F4E79"), ("Heading 2", 12, "2F5597")]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = __import__("docx").shared.RGBColor(31, 78, 121)
    p.paragraph_format.space_after = Pt(4)
    p = doc.add_paragraph("Malanga Community Welfare Scheme")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.paragraph_format.space_after = Pt(14)


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def registration_document():
    doc = Document()
    setup(doc, "Membership Registration and Approval Process")
    doc.add_paragraph("Approved working draft for the Malanga Community Welfare Scheme. This document defines the member journey, eligibility rules, application form, and information to be stored.")

    doc.add_heading("1. Membership Registration and Approval Process", level=1)
    steps = [
        ("Online Application", "A prospective member submits the online application form and provides all required personal details and supporting information."),
        ("Application Acknowledgement", "Immediately after submission, the applicant receives an automated notification confirming receipt and stating that the application is awaiting Welfare Committee review."),
        ("Committee Review", "The Welfare Committee reviews the application for eligibility, completeness, and compliance with membership requirements. If rejected, the applicant is notified. If approved, the applicant receives instructions to pay the registration fee."),
        ("Registration Fee Payment", "The approved applicant pays using a unique payment account or reference code. The code may be the registered Safaricom phone number or a system-generated identifier."),
        ("Membership Activation", "After payment is received and verified, the Committee receives a payment notification, the application is activated, a unique Membership Number is generated, and the new member receives confirmation with their membership details."),
        ("Probation Period", "The probation period begins on the date of successful activation and payment confirmation. During probation, the member is registered but is not yet eligible for full welfare benefits."),
        ("Full Membership", "After successful completion of probation, the system automatically changes the member to Full Membership. The member becomes eligible for benefits and services subject to the Welfare Constitution and Regulations."),
    ]
    for title, body in steps:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)

    doc.add_heading("2. Age Eligibility and Probation Policy", level=1)
    doc.add_paragraph("Age is calculated from the applicant’s date of birth as at the membership activation date, after payment verification.")
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, text in enumerate(["Age at Activation", "Admission Decision", "Probation Period"]):
        set_cell_text(table.rows[0].cells[i], text, True); shade(table.rows[0].cells[i], "D9EAF7")
    rows = [("Up to 50 years", "Eligible", "3 months"), ("51–75 years", "Eligible", "6 months"), ("Above 75 years", "Not eligible for admission", "Not applicable")]
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text)
    doc.add_paragraph("Applicants above 75 years must not be admitted under this policy. The Committee may reject an application for any other constitutional or verification reason.")

    doc.add_heading("3. Online Self-Registration and Application Form", level=1)
    sections = [
        ("A. Personal Information", ["1. Full Name (as it appears on National ID)", "2. National ID Number", "3. Date of Birth", "4. Gender", "5. Phone Number (Safaricom)", "6. Alternative Phone Number", "7. Email Address (Optional)"]),
        ("B. Residence Information", ["8. Are you a resident of Malanga? (Yes / No)", "9. If Yes, select your village:"] + [f"   {i}. {v}" for i, v in enumerate(VILLAGES, 1)] + ["10. If No, indicate your current residence/location (example: Malindi – Shella)"]),
        ("C. Dependants Information", ["11. Do you have dependants? (Yes / No)", "12. Dependant Details — repeatable for each dependant:", "   • Full Name", "   • Relationship to Applicant (Spouse, Child, Parent, Sibling, Other)", "   • Date of Birth (Optional)"]),
        ("D. Next of Kin Information", ["13. Next of Kin Full Name", "14. Relationship to Applicant", "15. Next of Kin Phone Number"]),
        ("E. Supporting Documents", ["16. Passport-size Photo Upload (Optional)"]),
        ("F. Declaration and Confirmation", ["17. Declaration: I declare that the information provided in this application is true and accurate to the best of my knowledge.", "18. Applicant Confirmation: Checkbox — I agree to the declaration above.", "19. Date of Application: Automatically captured by the system."]),
    ]
    for heading, items in sections:
        doc.add_heading(heading, level=2)
        add_bullets(doc, items)

    doc.add_heading("4. Data Summary to Be Stored", level=1)
    add_bullets(doc, ["Full Name", "National ID Number", "Date of Birth", "Gender", "Phone Number", "Alternative Phone Number", "Email Address (optional)", "Residence Status: Malanga Resident / Non-Resident", "Village or Current Location", "Dependants’ names, relationships, and optional dates of birth", "Next of Kin details", "Passport Photo (optional)", "Application Date", "Declaration Acceptance Status", "Application status and Committee decision", "Payment reference, payment status, and payment date", "Membership Number, activation date, probation end date, and membership status"])

    doc.add_heading("5. Required Notifications", level=1)
    add_bullets(doc, ["Application received and awaiting review", "Application approved and registration fee payment requested", "Application rejected, including a recorded reason where appropriate", "Payment received and under verification", "Membership activated with Membership Number and probation end date", "Probation completed and Full Membership attained"])
    doc.add_heading("6. Administrative Controls", level=1)
    add_bullets(doc, ["Only authorised Welfare Committee users may approve or reject applications.", "All decisions and status changes must be recorded with the responsible user and timestamp.", "Payment confirmation must be based on a trusted payment notification or reconciliation process.", "The system must prevent duplicate National ID numbers and duplicate active phone numbers unless authorised.", "Personal information and uploaded photos must be protected through role-based access and secure storage."])
    return doc


def implementation_document():
    doc = Document()
    setup(doc, "Membership Registration and Approval Implementation Plan")
    doc.add_paragraph("A phased implementation plan for digitising the Malanga Community Welfare membership application, review, payment, activation, probation, and full-membership process.")

    doc.add_heading("1. Objectives", level=1)
    add_bullets(doc, ["Provide a simple online self-registration process for applicants inside and outside Malanga.", "Give the Welfare Committee a controlled review and approval workflow.", "Automate payment references, payment verification, membership-number generation, notifications, and probation tracking.", "Maintain accurate, secure, and auditable member records."])

    doc.add_heading("2. Scope and Core Modules", level=1)
    modules = [("Applicant portal", "Application form, document upload, confirmation, and application-status tracking."), ("Committee dashboard", "Application queue, search, verification, approve/reject actions, notes, and audit history."), ("Payment module", "Unique reference, payment instructions, callback/reconciliation, and payment status."), ("Membership register", "Membership Number, profile, dependants, next of kin, status, and downloadable records."), ("Probation engine", "Age-band rules, probation start/end dates, reminders, and automatic full-membership transition."), ("Notifications", "SMS/email templates for each application, approval, payment, activation, rejection, and probation event."), ("Administration and reporting", "User roles, villages, configuration, reports, exports, and audit logs.")]
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(["Module", "Minimum capability"]): set_cell_text(table.rows[0].cells[i], text, True); shade(table.rows[0].cells[i], "D9EAF7")
    for row in modules:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text)

    doc.add_heading("3. Business Rules to Configure", level=1)
    add_numbered(doc, ["Applicants may be Malanga residents or non-residents.", "Malanga residents select one of the 14 configured villages; non-residents enter their current location.", "Applicants above 75 years are not eligible for admission.", "Applicants aged up to 50 years receive a 3-month probation period.", "Applicants aged 51–75 years receive a 6-month probation period.", "Probation starts only after payment is verified and membership is activated.", "Full Membership is granted automatically when the probation end date is reached, subject to any recorded suspension or constitutional restriction.", "Membership Number is generated once, is unique, and is never reassigned."])

    doc.add_heading("4. Delivery Phases", level=1)
    phases = [("Phase 1 — Requirements and governance", "Confirm Constitution requirements, fields, approval authority, rejection reasons, payment provider, notification channels, privacy requirements, and reporting needs.", "Approved requirements and decision matrix"), ("Phase 2 — User experience and data design", "Design applicant form, committee dashboard, statuses, database fields, village list, membership-number format, and notification templates.", "Signed-off screens and data model"), ("Phase 3 — Build the application workflow", "Implement registration, validation, uploads, acknowledgement, committee queue, approval/rejection, and audit trail.", "Working application workflow"), ("Phase 4 — Integrate payments and notifications", "Configure payment reference, payment callback or reconciliation, SMS/email messages, retries, and exception handling.", "Verified payment-to-activation flow"), ("Phase 5 — Probation and membership register", "Implement age calculation, probation dates, reminders, automatic full-membership transition, member profile, and reports.", "Complete membership lifecycle"), ("Phase 6 — Testing and data protection", "Run functional, security, usability, payment, notification, and recovery tests; fix defects and approve release.", "User acceptance sign-off"), ("Phase 7 — Pilot and rollout", "Pilot with a small Committee group and selected applicants, monitor results, train users, then launch to all applicants.", "Production launch and support plan")]
    table = doc.add_table(rows=1, cols=3); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(["Phase", "Activities", "Deliverable"]): set_cell_text(table.rows[0].cells[i], text, True); shade(table.rows[0].cells[i], "D9EAF7")
    for row in phases:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text)

    doc.add_heading("5. Recommended Application Statuses", level=1)
    add_bullets(doc, ["Draft", "Submitted — Awaiting Review", "Under Review", "Approved — Payment Pending", "Payment Pending Verification", "Active — Probation", "Active — Full Member", "Rejected", "Suspended", "Withdrawn"])

    doc.add_heading("6. Roles and Responsibilities", level=1)
    roles = [("Applicant", "Submit accurate information, accept the declaration, pay the fee after approval, and respond to requests."), ("Welfare Committee Reviewer", "Check eligibility, completeness, supporting information, and make a documented recommendation."), ("Welfare Committee Approver", "Approve or reject applications and authorise exceptional decisions."), ("Treasurer/Finance Officer", "Monitor payment notifications, reconcile exceptions, and confirm payment records."), ("System Administrator", "Manage users, roles, villages, templates, configuration, backups, and audit access."), ("System/Implementation Team", "Build, test, deploy, secure, monitor, and support the system.")]
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(["Role", "Responsibilities"]): set_cell_text(table.rows[0].cells[i], text, True); shade(table.rows[0].cells[i], "D9EAF7")
    for row in roles:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text)

    doc.add_heading("7. Testing Checklist", level=1)
    add_bullets(doc, ["Required-field and National ID validation", "Duplicate National ID and phone-number handling", "Resident village selection and non-resident location entry", "Repeatable dependant records", "Optional photo upload and file-type/size restrictions", "Acknowledgement, approval, rejection, payment, activation, and probation notifications", "Correct age calculation at activation", "Correct 3-month and 6-month probation dates", "Automatic rejection for applicants above 75", "Duplicate payment and delayed payment handling", "Unique Membership Number generation", "Role permissions and audit logs", "Data backup, recovery, and secure access", "Mobile-phone usability and low-bandwidth behaviour"])

    doc.add_heading("8. Launch Readiness Checklist", level=1)
    add_bullets(doc, ["Constitution and fee rules approved by the Welfare Committee", "Payment account and reconciliation owner confirmed", "SMS/email sender and message templates configured", "Committee users created with least-privilege access", "Village list and membership-number format confirmed", "Privacy notice and consent wording approved", "Support contact and escalation process published", "Pilot results reviewed and outstanding critical defects closed", "Backup and incident-response procedures tested"])

    doc.add_heading("9. Key Risks and Controls", level=1)
    risks = [("Incorrect or fraudulent details", "Require declaration, National ID review, Committee verification, and audit trail."), ("Payment received but membership not activated", "Use payment callbacks plus daily reconciliation and exception alerts."), ("Unauthorised access to personal data", "Use role-based access, secure storage, strong passwords, and access logs."), ("Applicants misunderstand probation", "Display probation dates and benefit eligibility clearly in approval and activation messages."), ("Connectivity or SMS failure", "Keep status visible in the portal and provide retry queues/manual contact procedures.")]
    table = doc.add_table(rows=1, cols=2); table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(["Risk", "Control"]): set_cell_text(table.rows[0].cells[i], text, True); shade(table.rows[0].cells[i], "D9EAF7")
    for row in risks:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text)

    doc.add_heading("10. Immediate Next Actions", level=1)
    add_numbered(doc, ["Approve the age and probation policy: up to 50 years = 3 months; 51–75 years = 6 months; above 75 years = no admission.", "Confirm the registration fee, payment provider, and whether the phone number will be the payment reference.", "Confirm the Welfare Committee approval quorum and authorised users.", "Approve the final form fields, declaration, privacy wording, and notification messages.", "Select the implementation team and agree the pilot date."])
    return doc


if __name__ == "__main__":
    registration_document().save("Malanga_Community_Welfare_Membership_Process.docx")
    implementation_document().save("Malanga_Community_Welfare_Implementation_Plan.docx")
    print("Created both DOCX files.")
